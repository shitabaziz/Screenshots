"""
xlsx builder for JobScholar Bangla GK MCQ sheets.

Same template as jobscholar_gk_builder.py (sheet "Questions", 7 columns, Arial,
wrap + top alignment, fixed column widths, "Live MCQ" sentence stripping and the
standard validation report), but it writes to memory so the web app can send
the file straight to the browser.
"""
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

HEADER = ["question", "option_a", "option_b", "option_c", "option_d", "correct_key", "explanation"]
WIDTHS = {"A": 50, "B": 22, "C": 22, "D": 22, "E": 22, "F": 12, "G": 85}

_LIVE_RE = re.compile(r'(?:^|(?<=[।.!?]\s))[^।.!?]*Live\s*MCQ[^।.!?]*[।.!?]', re.IGNORECASE)


def strip_live(text):
    """Remove 'Live MCQ' promotional sentences (full-sentence removal only)."""
    text = _LIVE_RE.sub('', text)
    return re.sub(r'[ \t]{2,}', ' ', text).strip()


def build_bytes(rows):
    """rows: [question, a, b, c, d, key, explanation] lists. Returns xlsx bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(HEADER)
    for c in range(1, 8):
        ws.cell(row=1, column=c).font = Font(name="Arial", bold=True)

    for q, a, b, c_, d_, key, exp in rows:
        ws.append([q, a, b, c_, d_, key, strip_live(exp)])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def validate_bytes(data):
    """Standard validation pipeline, run on the built file."""
    ws = load_workbook(io.BytesIO(data)).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    report = {
        "row_count": len(rows),
        "empty_field_rows": sorted({
            i + 2 for i, row in enumerate(rows)
            for v in row if v is None or (isinstance(v, str) and v.strip() == "")
        }),
        "bad_keys": [i + 2 for i, row in enumerate(rows) if row[5] not in ("a", "b", "c", "d")],
        "live_mcq_rows": [i + 2 for i, row in enumerate(rows) if "live mcq" in str(row[6]).lower()],
    }
    seen, dups = {}, []
    for i, row in enumerate(rows):
        key = tuple(row[:5])
        if key in seen:
            dups.append((seen[key], i + 2))
        else:
            seen[key] = i + 2
    report["duplicate_row_pairs"] = dups
    return report
