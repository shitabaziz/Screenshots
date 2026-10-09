"""
Reusable JobScholar Bangla GK MCQ xlsx builder.
(For the "general knowledge screenshot -> xlsx" project — NOT the Charyapada project,
which uses a different builder/ruleset with rewritten explanations + copyright line.)

Usage:
    from jobscholar_gk_builder import build, D

    rows = [
        ["question", "opt_a", "opt_b", "opt_c", "opt_d", "correct_key", "explanation_text" + D],
        ...
    ]
    build(rows, "topic_filename_slug")
    # saves to /mnt/user-data/outputs/topic_filename_slug.xlsx
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import os
import re

# Adjust this per-topic to reflect the actual citation source type seen in the
# screenshots (e.g. official org websites, Bangladesh Economic Survey, Wikipedia-style
# encyclopedic references, etc.) before building a new topic.
D = "\n\nতথ্যসূত্র: সংশ্লিষ্ট সংস্থার অফিসিয়াল ওয়েবসাইট।"


def strip_live(text):
    """Remove 'Live MCQ' promotional sentences (full-sentence removal only)."""
    pattern = re.compile(r'(?:^|(?<=[।.!?]\s))[^।.!?]*Live\s*MCQ[^।.!?]*[।.!?]', re.IGNORECASE)
    return pattern.sub('', text).strip()


def build(rows, filename_slug):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(["question", "option_a", "option_b", "option_c", "option_d", "correct_key", "explanation"])
    for c in range(1, 8):
        ws.cell(row=1, column=c).font = Font(name="Arial", bold=True)

    for row in rows:
        q, a, b, c_, d_, key, exp = row
        exp = strip_live(exp)
        ws.append([q, a, b, c_, d_, key, exp])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col, w in {"A": 50, "B": 22, "C": 22, "D": 22, "E": 22, "F": 12, "G": 85}.items():
        ws.column_dimensions[col].width = w

    os.makedirs("/mnt/user-data/outputs", exist_ok=True)
    out = f"/mnt/user-data/outputs/{filename_slug}.xlsx"
    wb.save(out)
    print(f"Saved: {out} with {len(rows)} rows")
    return out


def validate(path):
    """Run the standard validation pipeline used every time a file is built."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    report = {}
    report["row_count"] = len(rows)
    report["empty_field_rows"] = [
        i + 2 for i, row in enumerate(rows)
        for v in row if v is None or (isinstance(v, str) and v.strip() == "")
    ]
    report["bad_keys"] = [i + 2 for i, row in enumerate(rows) if row[5] not in ("a", "b", "c", "d")]
    report["live_mcq_rows"] = [i + 2 for i, row in enumerate(rows) if "live mcq" in str(row[6]).lower()]

    seen = {}
    dups = []
    for i, row in enumerate(rows):
        key = tuple(row[:5])
        if key in seen:
            dups.append((seen[key], i + 2))
        else:
            seen[key] = i + 2
    report["duplicate_row_pairs"] = dups

    for k, v in report.items():
        print(f"{k}: {v}")
    return report
